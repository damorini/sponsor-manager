"""Import della rubrica condivisa (Excel o CSV).

Colonne: company, referente, email, ruolo, tel, interessi (separati da ;).
Opzionali nome e cognome: se presenti vincono su 'referente'.
Due passaggi: analizza() non scrive nulla e dice cosa succederebbe riga per
riga; applica() rifa' l'analisi sui dati aggiornati e scrive SOLO le righe
senza errori. Mai unire aziende o creare aree in silenzio."""
import csv
import io
import re
import logging
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone

logger = logging.getLogger(__name__)

COLONNE_OBBLIGATORIE =['company', 'email']

# Il template che si scarica dall'utility (Importa contatti) usa nomi di
# colonna diversi dai nostri. Accettarli come sinonimi evita che chi scarica
# quel file si senta dire "colonne mancanti" proprio da questa pagina.
SINONIMI = {
    'sponsor_ragione_sociale': 'company',
    'azienda': 'company',
    'ragione_sociale': 'company',
    'ruolo_aziendale': 'ruolo',
    'telefono': 'tel',
    'aree_interesse': 'interessi',
    'aree_di_interesse': 'interessi',
}
SUFFISSI_SOCIETARI = r'\b(s\.?r\.?l\.?s?|s\.?p\.?a\.?|s\.?n\.?c\.?|s\.?a\.?s\.?|ltd|gmbh|inc|unipersonale)\b'


@dataclass
class RigaImport:
    numero: int
    azienda: str
    nome: str
    cognome: str
    email: str
    ruolo: str
    telefono: str
    aree: list = field(default_factory=list)
    esito: str = 'nuovo'
    messaggi: list = field(default_factory=list)
    nuova_azienda: bool = False
    sponsor_id: object = None

    def errore(self, msg):
        self.esito = 'errore'
        self.messaggi.append(msg)


def leggi_file(f):
    """Ritorna le righe non vuote come dict {intestazione minuscola: testo},
    con la chiave '_riga' = numero di riga nel file (intestazione = 1)."""
    nome = (getattr(f, 'name', '') or '').lower()
    dati = f.read()
    if nome.endswith('.xls'):
        raise ValueError("Formato .xls non supportato: salva il file come .xlsx o .csv.")
    if nome.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(dati), read_only=True, data_only=True)
        except Exception as e:
            raise ValueError("Impossibile leggere il file Excel: è danneggiato o non è un .xlsx.") from e
        try:
            if wb.active is None:
                raise ValueError("Il file Excel non ha fogli.")
            tabella = [['' if v is None else str(v) for v in r]
                       for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        try:
            testo = dati.decode('utf-8-sig')
        except UnicodeDecodeError:
            testo = dati.decode('cp1252', errors='replace')
        primo = testo.splitlines()[0] if testo else ''
        sep = ';' if primo.count(';') >= primo.count(',') else ','
        tabella = list(csv.reader(io.StringIO(testo), delimiter=sep))
    if not tabella:
        raise ValueError("Il file è vuoto.")
    intest = [SINONIMI.get(h.strip().lower(), h.strip().lower())
              for h in tabella[0]]
    mancanti = [c for c in COLONNE_OBBLIGATORIE if c not in intest]
    if 'referente' not in intest and 'cognome' not in intest:
        mancanti.append('referente (oppure cognome)')
    if mancanti:
        raise ValueError("Colonne mancanti nell'intestazione: " + ', '.join(mancanti))
    righe = []
    for n, valori in enumerate(tabella[1:], start=2):
        if not any((v or '').strip() for v in valori):
            continue
        r = {h: (valori[i].strip() if i < len(valori) and valori[i] else '')
             for i, h in enumerate(intest) if h}
        r['_riga'] = n
        righe.append(r)
    if not righe:
        # solo intestazione: senza questo l'anteprima resterebbe muta
        raise ValueError("Nessuna riga trovata nel file.")
    return righe


def normalizza_azienda(nome):
    s = (nome or '').casefold()
    s = re.sub(SUFFISSI_SOCIETARI, ' ', s)
    s = re.sub(r'[^\w]+', ' ', s)
    return ' '.join(s.split())


def dividi_referente(testo):
    """'Mario Rossi' -> ('Mario', 'Rossi'): ultima parola = cognome, come
    Contact.save() quando c'e' solo il nome completo."""
    parti = (testo or '').split()
    if not parti:
        return '', ''
    return ' '.join(parti[:-1]), parti[-1]


def _limiti():
    """(campo della riga, etichetta, max_length) letti dai modelli."""
    from sponsors.models import Contact, Sponsor
    voci = [('azienda', 'Azienda', Sponsor, 'legal_name'),
            ('nome', 'Nome', Contact, 'first_name'),
            ('cognome', 'Cognome', Contact, 'last_name'),
            ('email', 'Email', Contact, 'email'),
            ('ruolo', 'Ruolo', Contact, 'job_title'),
            ('telefono', 'Telefono', Contact, 'phone')]
    return [(campo, et, m._meta.get_field(nome).max_length) for campo, et, m, nome in voci]


def analizza(righe):
    from sponsors.models import Contact, InterestArea, Sponsor, SuppressedEmail
    limiti = _limiti()
    email_viste = {}
    nuove_grafie = {}

    aree_attive = {a.name.casefold(): a for a in InterestArea.objects.filter(is_active=True)}
    per_nome, per_norm = {}, {}
    for s in Sponsor.objects.all():
        for n in (s.legal_name, s.display_name):
            if n:
                per_nome.setdefault(n.strip().casefold(), s)
                per_norm.setdefault(normalizza_azienda(n), s)
    esclusi = dict(SuppressedEmail.objects.values_list('email', 'reason'))
    visti = {}
    out = []
    for r in righe:
        nome, cognome = r.get('nome', ''), r.get('cognome', '')
        if not (nome or cognome):
            nome, cognome = dividi_referente(r.get('referente', ''))
        riga = RigaImport(
            numero=r.get('_riga', 0), azienda=r.get('company', '').strip(),
            nome=nome.strip(), cognome=cognome.strip(),
            email=r.get('email', '').strip(), ruolo=r.get('ruolo', ''),
            telefono=r.get('tel', ''))
        email_l = riga.email.lower()

        if not riga.azienda:
            riga.errore("Manca l'azienda (colonna company).")
        if not riga.cognome:
            riga.errore("Manca il referente.")
        try:
            validate_email(riga.email)
        except ValidationError:
            riga.errore(f"Email mancante o non valida: «{riga.email}».")

        for campo, etichetta, massimo in limiti:
            valore = getattr(riga, campo)
            if massimo and len(valore) > massimo:
                riga.errore(f"{etichetta} troppo lungo ({len(valore)} caratteri, massimo {massimo}).")

        for nome_area in [a.strip() for a in re.split(r'[;,]', r.get('interessi', '')) if a.strip()]:
            area = aree_attive.get(nome_area.casefold())
            if area is None:
                riga.errore(f"Area di interesse sconosciuta: «{nome_area}». "
                            "Creala prima in Aree di interesse o correggi il file.")
            elif area.name not in riga.aree:
                riga.aree.append(area.name)

        motivo = esclusi.get(email_l)
        if motivo == SuppressedEmail.Reason.ANONYMIZED:
            riga.errore("Dati cancellati su richiesta dell'interessato (GDPR): non va reinserito.")
        elif motivo == SuppressedEmail.Reason.UNSUBSCRIBED:
            riga.messaggi.append("Si è disiscritto dal marketing: entra in rubrica ma non riceverà campagne.")

        sponsor = per_nome.get(riga.azienda.casefold())
        norm = normalizza_azienda(riga.azienda)
        if sponsor is not None:
            riga.sponsor_id = sponsor.pk
        elif riga.azienda and norm:
            simile = per_norm.get(norm)
            if simile is not None:
                riga.errore(f"Esiste già un'azienda simile: «{simile.legal_name}». "
                            "Se è la stessa, scrivi nel file il nome esatto.")
            else:
                riga.nuova_azienda = True
                riga.messaggi.append(f"Nuova azienda: verrà creata «{riga.azienda}».")
                grafia = riga.azienda.strip().casefold()
                prima = nuove_grafie.setdefault(norm, (grafia, riga.azienda, riga.numero))
                if prima[0] != grafia:
                    riga.errore(f"Nello stesso file l'azienda compare anche come «{prima[1]}» "
                                f"(riga {prima[2]}): usa lo stesso nome.")

        if email_l:
            altra = email_viste.setdefault(email_l, (norm, riga.numero))
            if altra[0] != norm:
                riga.errore("Stessa email già usata per un'altra azienda "
                            f"alla riga {altra[1]} del file.")

        chiave = (norm, email_l)
        if chiave in visti:
            riga.errore(f"Stessa azienda ed email già alla riga {visti[chiave]} del file.")
        else:
            visti[chiave] = riga.numero

        if riga.esito != 'errore' and sponsor is not None:
            esistente = Contact.objects.filter(sponsor=sponsor, email__iexact=riga.email).first()
            if esistente and esistente.left_company_at:
                riga.errore("Il contatto risulta uscito da questa azienda.")
            elif esistente:
                riga.esito = 'aggiorna'
        if riga.esito == 'nuovo':
            prova = Contact(sponsor=sponsor, first_name=riga.nome,
                            last_name=riga.cognome, email=riga.email)
            try:
                prova.clean()
            except ValidationError as e:
                for msgs in e.message_dict.values():
                    for m in msgs:
                        riga.errore(m)
        out.append(riga)
    return out


def applica(righe):
    from sponsors.models import Contact, InterestArea, Sponsor

    esito = {'creati': 0, 'aggiornati': 0, 'aziende_create': 0, 'scartati': 0}
    aree = {a.name: a for a in InterestArea.objects.all()}
    nuove_aziende = {}
    adesso = timezone.now()
    for r in analizza(righe):
        if r.esito == 'errore':
            esito['scartati'] += 1
            continue
        chiave = normalizza_azienda(r.azienda)
        azienda_creata = False
        try:
            with transaction.atomic():
                if r.nuova_azienda:
                    sponsor = nuove_aziende.get(chiave)
                    if sponsor is None:
                        sponsor = Sponsor.objects.create(legal_name=r.azienda, address_country='IT')
                        nuove_aziende[chiave] = sponsor
                        azienda_creata = True
                else:
                    sponsor = Sponsor.objects.get(pk=r.sponsor_id)
                contatto = Contact.objects.filter(sponsor=sponsor, email__iexact=r.email).first()
                nuovo = contatto is None
                if nuovo:
                    contatto = Contact(sponsor=sponsor, email=r.email,
                                       first_name=r.nome, last_name=r.cognome)
                if r.telefono:
                    contatto.phone = r.telefono
                if r.ruolo:
                    contatto.job_title = r.ruolo
                if not contatto.marketing_consent:
                    contatto.marketing_consent = True
                    contatto.marketing_consent_at = adesso
                contatto.save()
                contatto.interest_areas.add(*[aree[n] for n in r.aree])
        except Exception:
            logger.exception("Import rubrica: riga %s scartata (%s)", r.numero, r.email)
            if azienda_creata:
                # la creazione e' stata annullata col rollback della riga
                nuove_aziende.pop(chiave, None)
            esito['scartati'] += 1
            continue
        esito['creati' if nuovo else 'aggiornati'] += 1
        if azienda_creata:
            esito['aziende_create'] += 1
    return esito
