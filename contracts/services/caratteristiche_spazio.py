"""Riepilogo tecnico dello spazio assegnato allo Sponsor (tipologia, misure,
altezza massima di allestimento, allaccio elettrico, acqua, internet, accesso)
e le «Indicazioni specifiche» dello stand. Usato nella Domanda di ammissione
(Allegato 1), nel Regolamento tecnico (Allegato 2) e nella mail del PASS.
Le dotazioni assenti (es. allaccio idrico «no») non si indicano."""
from venues.models import misura

ETICHETTE_ACCESSO = ('Accesso al padiglione n°', 'Hall access no.')


def _stand(contract):
    if contract.stand_id:
        return [contract.stand]
    if contract.stand_block_id:
        return list(contract.stand_block.stands.all().order_by('code'))
    return []


def voci_stand(s, en=False):
    """[(etichetta, valore)] delle caratteristiche PRESENTI di uno stand."""
    voci = []
    if s.stand_type:
        voci.append(('Type' if en else 'Tipologia', s.stand_type))
    if s.dimensioni_testo:
        voci.append(('Size' if en else 'Misure', s.dimensioni_testo))
    if s.max_height_meters:
        voci.append(('Maximum set-up height' if en else 'Altezza max di allestimento',
                     f"{misura(s.max_height_meters)} m"))
    if s.power_kw:
        voci.append(('Electrical connection' if en else 'Allaccio elettrico',
                     f"{misura(s.power_kw)} kW"))
    elif s.has_power:
        voci.append(('Electrical connection' if en else 'Allaccio elettrico',
                     'yes' if en else 'sì'))
    if s.has_water:
        voci.append(('Water connection' if en else 'Allaccio idrico', 'yes' if en else 'sì'))
    if s.has_internet:
        voci.append(('Internet', 'yes' if en else 'sì'))
    if (s.access_door or '').strip():
        voci.append((ETICHETTE_ACCESSO[1] if en else ETICHETTE_ACCESSO[0],
                     s.access_door.strip()))
    return voci


def schede(contract, con_accesso=True):
    """Una scheda per stand (piu' d'una per un blocco):
    [{'codice', 'voci', 'indicazioni'}]."""
    en = (contract.language or 'it') == 'en'
    out = []
    for s in _stand(contract):
        voci = voci_stand(s, en)
        if not con_accesso:
            voci = [v for v in voci if v[0] not in ETICHETTE_ACCESSO]
        out.append({'codice': s.code, 'voci': voci,
                    'indicazioni': (s.caratteristiche or '').strip()})
    return [x for x in out if x['voci'] or x['indicazioni']]


# ---------------------------------------------------------------------------
# Tabellina Word (Domanda di ammissione e Regolamento tecnico)
# ---------------------------------------------------------------------------

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _run(par, testo, size, bold, font=None, colore=None):
    from docx.shared import Pt, RGBColor
    r = par.add_run(testo)
    r.font.size = Pt(size)
    r.bold = bold
    if font:
        r.font.name = font
    if colore:
        r.font.color.rgb = RGBColor.from_string(colore)
    return r


def _fondo(cella, colore):
    from docx.oxml import OxmlElement
    tcpr = cella._tc.get_or_add_tcPr()
    for vecchio in tcpr.findall(W + 'shd'):
        tcpr.remove(vecchio)
    shd = OxmlElement('w:shd')
    shd.set(W + 'val', 'clear')
    shd.set(W + 'color', 'auto')
    shd.set(W + 'fill', colore)
    tcpr.append(shd)


def _paragrafo_stretto(par, prima=1, dopo=1):
    from docx.shared import Pt
    par.paragraph_format.space_before = Pt(prima)
    par.paragraph_format.space_after = Pt(dopo)
    par.paragraph_format.line_spacing = 1.0


def tabella_docx(doc, scheda, en, larghezza_twip, size=9, font=None, titolo=True):
    """Tabella «RIEPILOGO TECNICO DELLO SPAZIO» (etichetta | valore), da
    posizionare dal chiamante: ritorna l'elemento <w:tbl>."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    voci = scheda['voci']
    righe = len(voci) + (1 if titolo else 0)
    t = doc.add_table(rows=max(righe, 1), cols=2)
    tbl = t._tbl
    tbl.getparent().remove(tbl)
    tblpr = tbl.tblPr
    bordi = OxmlElement('w:tblBorders')
    for lato in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement(f'w:{lato}')
        b.set(W + 'val', 'single')
        b.set(W + 'sz', '4')
        b.set(W + 'color', '9CA3AF')
        bordi.append(b)
    tblpr.append(bordi)
    for vecchio in tblpr.findall(W + 'tblW'):
        tblpr.remove(vecchio)
    tblw = OxmlElement('w:tblW')
    tblw.set(W + 'w', str(larghezza_twip))
    tblw.set(W + 'type', 'dxa')
    tblpr.append(tblw)
    quote = (int(larghezza_twip * 0.40), larghezza_twip - int(larghezza_twip * 0.40))
    for gc, q in zip(tbl.find(W + 'tblGrid').findall(W + 'gridCol'), quote):
        gc.set(W + 'w', str(q))
    for riga in t.rows:
        for cella, q in zip(riga.cells, quote):
            tcw = cella._tc.get_or_add_tcPr().get_or_add_tcW()
            tcw.set(W + 'w', str(q))
            tcw.set(W + 'type', 'dxa')
    i = 0
    if titolo:
        intest = t.rows[0].cells[0].merge(t.rows[0].cells[1])
        par = intest.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _paragrafo_stretto(par, 2, 2)
        testo = ('TECHNICAL SUMMARY OF THE SPACE' if en else
                 'RIEPILOGO TECNICO DELLO SPAZIO') + f" – STAND {scheda['codice']}"
        _run(par, testo, size, True, font)
        _fondo(intest, 'D9D9D9')
        i = 1
    for etichetta, valore in voci:
        c_et, c_val = t.rows[i].cells
        _paragrafo_stretto(c_et.paragraphs[0])
        _paragrafo_stretto(c_val.paragraphs[0])
        _run(c_et.paragraphs[0], etichetta, size, True, font)
        _fondo(c_et, 'F3F4F6')
        _run(c_val.paragraphs[0], str(valore), size, False, font)
        i += 1
    return tbl


def paragrafi_indicazioni(doc, dopo_el, scheda, en, size=9, font=None):
    """«Indicazioni specifiche:» e le righe del testo, subito dopo dopo_el.
    Ritorna l'ultimo elemento inserito."""
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    testo = scheda.get('indicazioni') or ''
    if not testo:
        return dopo_el
    righe = [r.strip() for r in testo.splitlines() if r.strip()]
    ultimo = dopo_el
    for n, riga in enumerate(righe):
        el = OxmlElement('w:p')
        ultimo.addnext(el)
        ultimo = el
        par = Paragraph(el, doc)
        _paragrafo_stretto(par, 4 if n == 0 else 0, 0)
        if n == 0:
            _run(par, ('Specific notes: ' if en else 'Indicazioni specifiche: '),
                 size, True, font)
        _run(par, riga, size, False, font)
    return ultimo


def larghezza_utile(doc):
    sp = doc.element.body.find(W + 'sectPr')
    pgsz, pgmar = sp.find(W + 'pgSz'), sp.find(W + 'pgMar')
    return (int(pgsz.get(W + 'w')) - int(pgmar.get(W + 'left'))
            - int(pgmar.get(W + 'right')))


def aggiungi_riepilogo_domanda(docx_path, contract):
    """Domanda di ammissione / Allegato 1: in fondo, dopo le firme, il
    RIEPILOGO TECNICO DELLO SPAZIO con le indicazioni specifiche."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Pt
    from docx.text.paragraph import Paragraph
    tutte = schede(contract)
    if not tutte:
        return False
    en = (contract.language or 'it') == 'en'
    doc = Document(str(docx_path))
    body = doc.element.body
    sectpr = body.find(W + 'sectPr')
    larghezza = larghezza_utile(doc)
    for scheda in tutte:
        stacco = OxmlElement('w:p')
        sectpr.addprevious(stacco)
        Paragraph(stacco, doc).paragraph_format.space_before = Pt(10)
        tbl = tabella_docx(doc, scheda, en, larghezza, size=9)
        sectpr.addprevious(tbl)
        paragrafi_indicazioni(doc, tbl, scheda, en, size=9)
    doc.save(str(docx_path))
    return True
