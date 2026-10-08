"""Regolamento tecnico: giorno e orari di disallestimento presi dalla tabella
dell'evento tramite segnaposto."""
from datetime import date, time

import pytest
from docx import Document

from contracts.models import Contract, ContractKind, ContractStatus
from events.models import Event, EventSetupDay


@pytest.fixture
def contratto(db, sponsor):
    ev = Event.objects.create(name={'it': 'Ev Dis'}, code='DIS',
                              start_date=date(2027, 2, 25), end_date=date(2027, 2, 27))
    return Contract.objects.create(sponsor=sponsor, event=ev, contract_kind=ContractKind.MAIN,
                                   status=ContractStatus.SIGNED, contract_number='DIS-001')


def _compila(contratto, tmp_path):
    from contracts.services.allegato2 import prepara_docx
    d0 = Document()
    d0.add_paragraph('Dalle ore {{ disallestimento_inizio }} fino alle ore '
                     '{{ disallestimento_fine }} del {{ disallestimento_giorno }}.')
    percorso = tmp_path / 'a2.docx'
    d0.save(str(percorso))
    prepara_docx(percorso, contratto)
    return '\n'.join(p.text for p in Document(str(percorso)).paragraphs)


def test_dalla_tabella(contratto, tmp_path):
    EventSetupDay.objects.create(event=contratto.event, kind='allestimento',
                                 date=date(2027, 2, 23), start_time=time(8), end_time=time(19))
    EventSetupDay.objects.create(event=contratto.event, kind='disallestimento',
                                 date=date(2027, 2, 27), start_time=time(17), end_time=time(23))
    assert 'Dalle ore 17:00 fino alle ore 23:00 del 27 febbraio 2027.' in _compila(
        contratto, tmp_path)


def test_senza_giorni(contratto, tmp_path):
    assert 'del giorno da comunicare.' in _compila(contratto, tmp_path)
