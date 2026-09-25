"""Allegato 2 al contratto (es. Regolamento tecnico): file Word caricato
sull'evento, accodato al contratto di sponsorizzazione dopo l'Allegato 1."""
import io
import re
from datetime import date
from decimal import Decimal

import pytest
from django.core.files.base import ContentFile
from docx import Document


def _docx(testo):
    d = Document()
    d.add_paragraph(testo)
    buf = io.BytesIO()
    d.save(buf)
    return ContentFile(buf.getvalue(), name='regolamento.docx')


@pytest.fixture
def contratto(db, sponsor, dati_firmatario_completi, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    from sponsors.models import Contact, ContactRole
    from events.models import Event, EventType
    from catalog.models import Service
    from contracts.models import Contract, ContractKind, ContractStatus, ContractLine
    Contact.objects.create(sponsor=sponsor, full_name='Firma Tario',
                           email='firma@test.it', roles=[ContactRole.OPERATIONAL],
                           is_signer=True, **dati_firmatario_completi)
    ev = Event.objects.create(name={'it': 'Ev All2'}, code='AL2',
                              event_type=EventType.NON_ECM,
                              start_date=date(2027, 2, 25), end_date=date(2027, 2, 27))
    c = Contract.objects.create(sponsor=sponsor, event=ev,
                                contract_kind=ContractKind.MAIN,
                                status=ContractStatus.SIGNED,
                                contract_number='AL2-27-001')
    s = Service.objects.create(event=ev, code='S', name={'it': 'Servizio'},
                               base_price=Decimal('1000'))
    ContractLine.objects.create(contract=c, service=s, quantity=1)
    return c


def _testo_pdf(document):
    from pathlib import Path
    from django.conf import settings
    from pypdf import PdfReader
    rel = document.storage_url.replace(settings.MEDIA_URL, '', 1)
    r = PdfReader(str(Path(settings.MEDIA_ROOT) / rel))
    return len(r.pages), '\n'.join(p.extract_text() or '' for p in r.pages)


def test_allegato_2_accodato(contratto):
    from contracts.services.pdf_generator import generate_sponsor_contract_pdf
    ev = contratto.event
    ev.contract_annex_enabled = True
    ev.contract_annex_file.save('regolamento.docx',
                                _docx('Articolo 1 del regolamento tecnico prova'))
    ev.save()
    doc = generate_sponsor_contract_pdf(contratto)
    if not str(doc.storage_url).endswith('.pdf'):
        pytest.skip('conversione PDF non disponibile in questo ambiente')
    pagine, testo = _testo_pdf(doc)
    assert 'ALLEGATO 2 – Regolamento tecnico' in testo
    assert 'Articolo 1 del regolamento tecnico prova' in testo
    assert testo.index('ALLEGATO 2') > testo.index('ALLEGATO 1')
    m = re.search(r'si compone di N\. (\d+) pagine', testo)
    assert m and int(m.group(1)) == pagine


def test_senza_spunta_niente_allegato_2(contratto):
    from contracts.services.pdf_generator import generate_sponsor_contract_pdf
    ev = contratto.event
    ev.contract_annex_file.save('regolamento.docx', _docx('Non deve comparire'))
    ev.save()
    doc = generate_sponsor_contract_pdf(contratto)
    if not str(doc.storage_url).endswith('.pdf'):
        pytest.skip('conversione PDF non disponibile in questo ambiente')
    _pagine, testo = _testo_pdf(doc)
    assert 'ALLEGATO 2' not in testo
    assert 'Non deve comparire' not in testo


def test_spunta_senza_file_non_si_salva(db):
    from events.admin import EventAdminForm
    from events.models import Event
    ev = Event.objects.create(name={'it': 'Ev X'}, code='X2',
                              start_date=date(2027, 2, 25), end_date=date(2027, 2, 27))
    form = EventAdminForm(instance=ev)
    form.cleaned_data = {'contract_annex_enabled': True}
    form._errors = {}
    form.clean()
    assert 'contract_annex_file' in form._errors


def _docx_regolamento():
    d = Document()
    d.add_paragraph('REGOLAMENTO TECNICO', style='Title')
    d.add_paragraph('Il congresso si svolge nei giorni {{ date_evento }}.')
    t = d.add_table(rows=2, cols=3)
    for c, testo in zip(t.rows[0].cells, ['Attività', 'Data', 'Orario']):
        c.text = testo
    for c in t.rows[1].cells:
        c.text = '[Da compilare]'
    d.sections[0].header.paragraphs[0].text = 'Pagina 1 di 9 (vecchia intestazione)'
    buf = io.BytesIO()
    d.save(buf)
    return ContentFile(buf.getvalue(), name='regolamento.docx')


def test_allegato_2_personalizzato(contratto, tmp_path):
    from datetime import time
    from events.models import EventSetupDay
    from venues.models import Stand
    from contracts.services.allegato2 import prepara_docx
    ev = contratto.event
    EventSetupDay.objects.create(event=ev, kind='allestimento', date=date(2027, 2, 24),
                                 start_time=time(8), end_time=time(20))
    EventSetupDay.objects.create(event=ev, kind='disallestimento', date=date(2027, 2, 27),
                                 start_time=time(17), end_time=time(23))
    contratto.stand = Stand.objects.create(event=ev, code='1-A',
                                           base_price=Decimal('1000'))
    contratto.save()
    percorso = tmp_path / 'a2.docx'
    percorso.write_bytes(_docx_regolamento().read())
    prepara_docx(percorso, contratto)

    d = Document(str(percorso))
    testo = '\n'.join(p.text for p in d.paragraphs)
    assert 'ALLEGATO 2 – Regolamento tecnico' in testo
    assert 'Azienda espositrice: ' + contratto.sponsor.legal_name in testo
    assert 'Stand n.: 1-A' in testo
    assert 'Contratto n.: AL2-27-001' in testo
    assert '25/02/2027 – 27/02/2027' in testo
    righe = [[c.text for c in r.cells] for r in d.tables[0].rows]
    assert righe[1:] == [['Allestimento', 'Mercoledì 24/02/2027', '08:00 – 20:00'],
                         ['Disallestimento', 'Sabato 27/02/2027', '17:00 – 23:00']]
    assert 'vecchia intestazione' not in ' '.join(
        p.text for p in d.sections[0].header.paragraphs)
    assert all(r.font.name == 'Arial' for p in d.paragraphs for r in p.runs if r.text)


def test_porta_accesso_e_note_giorni(contratto, tmp_path):
    from datetime import time
    from events.models import EventSetupDay
    from venues.models import Stand
    from contracts.services.allegato2 import prepara_docx
    ev = contratto.event
    EventSetupDay.objects.create(
        event=ev, kind='allestimento', date=date(2027, 2, 25),
        start_time=time(8), end_time=time(12),
        notes='NON PER ALLESTIMENTO - solo posizionamento materiale')
    EventSetupDay.objects.create(event=ev, kind='disallestimento', date=date(2027, 2, 27),
                                 start_time=time(17), end_time=time(23))
    contratto.stand = Stand.objects.create(event=ev, code='1-B', access_door='3',
                                           base_price=Decimal('1000'))
    contratto.save()
    percorso = tmp_path / 'a2.docx'
    percorso.write_bytes(_docx_regolamento().read())
    prepara_docx(percorso, contratto)

    d = Document(str(percorso))
    testo = '\n'.join(p.text for p in d.paragraphs)
    assert 'Accesso al padiglione n°: 3' in testo
    assert "tramite l'accesso n° 3." in testo
    righe = [[c.text for c in r.cells] for r in d.tables[0].rows]
    # riga del giorno, poi la sua nota su tutta la larghezza, poi il giorno dopo
    assert righe[1] == ['Allestimento', 'Giovedì 25/02/2027', '08:00 – 12:00']
    assert righe[2][0] == 'Nota: NON PER ALLESTIMENTO - solo posizionamento materiale'
    assert righe[3][0] == 'Disallestimento'


def test_senza_porta_niente_frase(contratto, tmp_path):
    from contracts.services.allegato2 import prepara_docx
    percorso = tmp_path / 'a2.docx'
    percorso.write_bytes(_docx_regolamento().read())
    prepara_docx(percorso, contratto)
    testo = '\n'.join(p.text for p in Document(str(percorso)).paragraphs)
    assert 'Accesso al padiglione' not in testo
    assert "tramite l'accesso" not in testo
