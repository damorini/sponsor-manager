"""RIEPILOGO TECNICO DELLO SPAZIO in tabellina (Domanda/Allegato 1 e
Regolamento tecnico/Allegato 2) e nel PASS; le dotazioni assenti non si
indicano; le «Indicazioni specifiche» dello stand vanno sotto la tabellina."""
from datetime import date
from decimal import Decimal

import pytest
from django.core import mail
from docx import Document

from contracts.models import Contract, ContractKind, ContractStatus
from events.models import Event
from venues.models import Stand


@pytest.fixture
def contratto(db, sponsor, contact):
    ev = Event.objects.create(name={'it': 'Ev Spazio'}, code='SPA',
                              start_date=date(2027, 2, 25), end_date=date(2027, 2, 27))
    st = Stand.objects.create(
        event=ev, code='1-A', base_price=Decimal('1000'), stand_type='Spazio Main-XL',
        width_meters=Decimal('6.00'), depth_meters=Decimal('4.00'),
        max_height_meters=Decimal('3.50'), has_power=True, power_kw=Decimal('3.00'),
        has_water=False, has_internet=False, access_door='3',
        caratteristiche='Pilastro sul lato corto\nNiente appendimenti a soffitto')
    return Contract.objects.create(sponsor=sponsor, event=ev, stand=st,
                                   contract_kind=ContractKind.MAIN,
                                   status=ContractStatus.SIGNED, contract_number='SPA-001')


def _tabella_riepilogo(d):
    for t in d.tables:
        if t.rows and 'RIEPILOGO TECNICO DELLO SPAZIO' in t.rows[0].cells[0].text:
            return {r.cells[0].text: r.cells[1].text for r in t.rows[1:]}, t
    return None, None


def test_voci_presenti_senza_i_no(contratto):
    from contracts.services.caratteristiche_spazio import schede
    s = schede(contratto)[0]
    voci = dict(s['voci'])
    assert voci == {'Tipologia': 'Spazio Main-XL', 'Misure': '6 × 4 m (24 m²)',
                    'Altezza max di allestimento': '3,5 m', 'Allaccio elettrico': '3 kW',
                    'Accesso al padiglione n°': '3'}
    assert s['indicazioni'].startswith('Pilastro')


def test_tabellina_nel_regolamento(contratto, tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    from contracts.services.allegato2 import prepara_docx
    d0 = Document()
    d0.add_paragraph('REGOLAMENTO TECNICO', style='Title')
    d0.add_paragraph('1. GLOSSARIO')
    percorso = tmp_path / 'a2.docx'
    d0.save(str(percorso))
    prepara_docx(percorso, contratto)
    d = Document(str(percorso))
    voci, t = _tabella_riepilogo(d)
    assert t is not None and 'STAND 1-A' in t.rows[0].cells[0].text
    assert voci['Misure'] == '6 × 4 m (24 m²)'
    assert voci['Allaccio elettrico'] == '3 kW'
    assert 'Allaccio idrico' not in voci and 'Internet' not in voci
    assert 'Accesso al padiglione n°' not in voci       # e' gia' nella tabellina sopra
    testo = '\n'.join(p.text for p in d.paragraphs)
    assert 'Indicazioni specifiche: Pilastro sul lato corto' in testo
    assert 'Niente appendimenti a soffitto' in testo


def test_tabellina_nella_domanda(contratto, tmp_path):
    from contracts.services.caratteristiche_spazio import aggiungi_riepilogo_domanda
    d0 = Document()
    d0.add_paragraph('DOMANDA DI AMMISSIONE')
    d0.add_paragraph('Firma')
    percorso = tmp_path / 'dom.docx'
    d0.save(str(percorso))
    assert aggiungi_riepilogo_domanda(percorso, contratto)
    d = Document(str(percorso))
    voci, t = _tabella_riepilogo(d)
    assert voci['Accesso al padiglione n°'] == '3'
    assert 'Allaccio idrico' not in voci
    testi = [p.text for p in d.paragraphs]
    assert testi.index('Firma') < testi.index(
        'Indicazioni specifiche: Pilastro sul lato corto')


def test_nel_pass(contratto):
    from contracts.services import pass_allestimento as pa
    pa.invia(contratto)
    html = mail.outbox[0].alternatives[0][0]
    assert 'Caratteristiche del vostro spazio' in html
    for atteso in ('6 × 4 m (24 m²)', '3,5 m', '3 kW', 'Indicazioni specifiche',
                   'Pilastro sul lato corto'):
        assert atteso in html, atteso
    assert 'Allaccio idrico' not in html


def test_excel_colonna_indicazioni(contratto):
    from catalog.utils.excel_template import build_template_stand_workbook
    intest = [c.value for c in build_template_stand_workbook().active[1]]
    assert 'indicazioni_specifiche' in intest
