"""Contratto di sponsorizzazione: nella tabella delle parti i valori (data di
nascita, email...) stanno su una riga sola."""
import io
from pathlib import Path

import pytest
from django.conf import settings
from docx import Document
from docx.shared import Pt

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
CARTELLA = Path(settings.BASE_DIR) / 'contracts' / 'templates_pdf'


@pytest.mark.parametrize('lingua', ['it', 'en'])
def test_colonne_ridistribuite(lingua):
    d = Document(str(CARTELLA / f'template_contratto_sponsor_non_ecm_{lingua}.docx'))
    grid = [int(g.get(W + 'w')) for g in
            d.tables[0]._tbl.find(W + 'tblGrid').findall(W + 'gridCol')]
    assert sum(grid) == 9747                  # stessa larghezza totale
    assert sum(grid[2:5]) >= 1300             # data di nascita
    assert sum(grid[4:8]) >= 3300             # email


def test_testo_troppo_lungo_rimpicciolito(tmp_path):
    from contracts.services.pdf_generator import _una_riga_tabella_parti
    d = Document()
    t = d.add_table(rows=1, cols=2)
    grid = t._tbl.find(W + 'tblGrid').findall(W + 'gridCol')
    grid[0].set(W + 'w', '1000')
    grid[1].set(W + 'w', '3000')
    corto = t.rows[0].cells[0].paragraphs[0].add_run('BO')
    lungo = t.rows[0].cells[1].paragraphs[0].add_run(
        'amministrazione.fornitori@azienda-lunghissima.it')
    corto.font.size = lungo.font.size = Pt(9)
    p = tmp_path / 'x.docx'
    d.save(str(p))
    _una_riga_tabella_parti(p)
    d2 = Document(str(p))
    celle = d2.tables[0].rows[0].cells
    assert celle[0].paragraphs[0].runs[0].font.size.pt == 9      # invariato
    assert 7 <= celle[1].paragraphs[0].runs[0].font.size.pt < 9  # ridotto
