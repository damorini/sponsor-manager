"""Accesso per montaggio e smontaggio: indirizzo sull'evento (diverso
dall'ingresso del congresso) nel Regolamento tecnico e nella mail del PASS."""
from datetime import date, time
from decimal import Decimal

import pytest
from django.core import mail
from docx import Document

from contracts.models import Contract, ContractKind, ContractStatus
from events.models import Event, EventSetupDay


@pytest.fixture
def contratto(db, sponsor, contact):
    ev = Event.objects.create(
        name={'it': 'Ev Montaggio'}, code='MON',
        start_date=date(2027, 2, 25), end_date=date(2027, 2, 27),
        montaggio_indirizzo='Via Calzoni 1/5 - Bologna',
        montaggio_dettagli='Mezzi non oltre 20 metri')
    EventSetupDay.objects.create(event=ev, kind='allestimento', date=date(2027, 2, 24),
                                 start_time=time(8), end_time=time(20))
    return Contract.objects.create(sponsor=sponsor, event=ev, contract_kind=ContractKind.MAIN,
                                   status=ContractStatus.SIGNED, contract_number='MON-001')


def _regolamento(tmp_path):
    d = Document()
    d.add_paragraph('REGOLAMENTO TECNICO', style='Title')
    d.add_paragraph('Ingresso automezzi da: {{ montaggio_indirizzo }}.')
    t = d.add_table(rows=2, cols=3)
    for c, testo in zip(t.rows[0].cells, ('Attività', 'Data', 'Orario')):
        c.text = testo
    percorso = tmp_path / 'a2.docx'
    d.save(str(percorso))
    return percorso


def test_regolamento_con_indirizzo_di_montaggio(contratto, tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    from contracts.services.allegato2 import prepara_docx
    from venues.models import Stand
    contratto.stand = Stand.objects.create(event=contratto.event, code='1-C',
                                           access_door='3', base_price=Decimal('1000'))
    contratto.save()
    percorso = _regolamento(tmp_path)
    prepara_docx(percorso, contratto)
    testo = '\n'.join(p.text for p in Document(str(percorso)).paragraphs)
    assert 'Ingresso automezzi da: Via Calzoni 1/5 - Bologna.' in testo
    assert ('Indirizzo di accesso per montaggio e smontaggio: '
            'Via Calzoni 1/5 - Bologna') in testo
    assert 'Mezzi non oltre 20 metri' in testo
    assert "tramite l'accesso n° 3." in testo


def test_regolamento_senza_indirizzo(contratto, tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    from contracts.services.allegato2 import prepara_docx
    Event.objects.filter(pk=contratto.event_id).update(montaggio_indirizzo='',
                                                       montaggio_dettagli='')
    contratto.refresh_from_db()
    percorso = _regolamento(tmp_path)
    prepara_docx(percorso, contratto)
    testo = '\n'.join(p.text for p in Document(str(percorso)).paragraphs)
    assert 'Ingresso automezzi da: da comunicare.' in testo
    assert 'Indirizzo di accesso per montaggio' not in testo


def test_pass_con_accesso_montaggio(contratto):
    from contracts.services import pass_allestimento as pa
    pa.invia(contratto)
    html = mail.outbox[0].alternatives[0][0]
    assert 'ACCESSO PER MONTAGGIO E SMONTAGGIO' in html
    assert 'Via Calzoni 1/5 - Bologna' in html
    assert 'Mezzi non oltre 20 metri' in html


def test_pass_senza_accesso_montaggio(contratto):
    from contracts.services import pass_allestimento as pa
    Event.objects.filter(pk=contratto.event_id).update(montaggio_indirizzo='')
    contratto.refresh_from_db()
    pa.invia(contratto)
    assert 'ACCESSO PER MONTAGGIO' not in mail.outbox[0].alternatives[0][0]
