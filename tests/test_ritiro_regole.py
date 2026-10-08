"""Evento: ritiro del materiale a fine evento e regole per montaggio/smontaggio
nella mail del PASS (e quindi nel PDF) e come segnaposto del Regolamento."""
import re
from datetime import date

import pytest
from django.core import mail
from docx import Document

from contracts.models import Contract, ContractKind, ContractStatus
from contracts.services import pass_allestimento as pa
from events.models import Event


@pytest.fixture
def contratto(db, sponsor, contact):
    ev = Event.objects.create(
        name={'it': 'Ev Ritiro'}, code='RIT',
        start_date=date(2027, 2, 25), end_date=date(2027, 2, 27),
        magazzino_ritiro='Ritiro entro lunedì ore 12\nColli etichettati',
        regole_montaggio='Vietato forare le pareti\nMoquette obbligatoria')
    return Contract.objects.create(sponsor=sponsor, event=ev, contract_kind=ContractKind.MAIN,
                                   status=ContractStatus.SIGNED, contract_number='RIT-1')


def _testo(html):
    return re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', html))


def test_nel_pass(contratto):
    pa.invia(contratto)
    t = _testo(mail.outbox[0].alternatives[0][0])
    assert "Regole per il montaggio e lo smontaggio dell'area espositiva" in t
    assert 'Vietato forare le pareti Moquette obbligatoria' in t
    assert 'Ritiro del materiale a fine evento' in t
    assert 'Ritiro entro lunedì ore 12 Colli etichettati' in t
    assert t.index('Regole per il montaggio') < t.index('Invio dei materiali') < t.index(
        'Ritiro del materiale a fine evento')


def test_sezioni_vuote_non_compaiono(contratto):
    Event.objects.filter(pk=contratto.event_id).update(magazzino_ritiro='', regole_montaggio='')
    contratto.refresh_from_db()
    pa.invia(contratto)
    t = _testo(mail.outbox[0].alternatives[0][0])
    assert 'Regole per il montaggio' not in t and 'Ritiro del materiale a fine' not in t


def test_segnaposto_regolamento(contratto, tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    from contracts.services.allegato2 import prepara_docx
    d0 = Document()
    d0.add_paragraph('{{ regole_montaggio }}')
    d0.add_paragraph('{{ magazzino_ritiro }}')
    percorso = tmp_path / 'a2.docx'
    d0.save(str(percorso))
    prepara_docx(percorso, contratto)
    testo = '\n'.join(p.text for p in Document(str(percorso)).paragraphs)
    assert 'Vietato forare le pareti' in testo and 'Ritiro entro lunedì ore 12' in testo
