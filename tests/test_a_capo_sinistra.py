"""Allegato 2: indirizzi su piu' righe non giustificati (righe «stirate»)."""
from datetime import date

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from contracts.models import Contract, ContractKind, ContractStatus
from events.models import Event


@pytest.fixture
def contratto(db, sponsor):
    ev = Event.objects.create(name={'it': 'Ev Capo'}, code='CAP',
                              start_date=date(2027, 2, 25), end_date=date(2027, 2, 27),
                              magazzino_indirizzo='MAGAZZINO\r\nINGRESSO PALCOSCENICO\r\nVia Calzoni 1/5')
    return Contract.objects.create(sponsor=sponsor, event=ev, contract_kind=ContractKind.MAIN,
                                   status=ContractStatus.SIGNED, contract_number='CAP-001')


def test_indirizzo_su_piu_righe_a_sinistra(contratto, tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    from contracts.services.allegato2 import prepara_docx
    d0 = Document()
    p = d0.add_paragraph('INDIRIZZO: {{ magazzino_indirizzo }}')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    q = d0.add_paragraph('Testo normale giustificato.')
    q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    percorso = tmp_path / 'a2.docx'
    d0.save(str(percorso))
    prepara_docx(percorso, contratto)
    d = Document(str(percorso))
    ind = next(x for x in d.paragraphs if x.text.startswith('INDIRIZZO'))
    assert 'INGRESSO PALCOSCENICO' in ind.text
    assert ind.alignment == WD_ALIGN_PARAGRAPH.LEFT
    norm = next(x for x in d.paragraphs if x.text.startswith('Testo normale'))
    assert norm.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
