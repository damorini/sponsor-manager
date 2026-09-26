"""Contratto: art. 12 «Esonero di responsabilità ed obblighi assicurativi»
dopo l'art. 11; modifiche e legge applicabile diventano 13 e 14."""
from pathlib import Path

import pytest
from django.conf import settings
from docx import Document

CARTELLA = Path(settings.BASE_DIR) / 'contracts' / 'templates_pdf'


@pytest.mark.parametrize('lingua,titolo,assicurazione,voce_1341,risolutiva', [
    ('it', '12. Esonero di responsabilità ed obblighi assicurativi', 'polizza assicurativa',
     '12. Esonero di responsabilità ed obblighi assicurativi; 13. Modifiche o variazioni; '
     '14. Legge applicabile e foro competente.', '9, 10, 11, 12, 13, 14 '),
    ('en', '12. Exemption from liability and insurance obligations', 'insurance policy',
     '12. Exemption from liability and insurance obligations; 13. Amendments or '
     'variations; 14. Governing law and jurisdiction.', '11, 12, 13 and 14 '),
])
def test_art_12(lingua, titolo, assicurazione, voce_1341, risolutiva):
    t = [p.text.strip() for p in Document(
        str(CARTELLA / f'template_contratto_sponsor_non_ecm_{lingua}.docx')).paragraphs]
    i = t.index(titolo)
    assert t[i + 1].startswith('12.1 ')
    assert sum(x.startswith('– ') for x in t[i + 2:i + 5]) == 3
    assert any(x.startswith('12.3 ') and assicurazione in x for x in t)
    assert any(x.startswith('12.4 ') for x in t)
    inizi = [x.split(' ')[0] for x in t]
    assert '13.1' in inizi and '14.1' in inizi and inizi.count('12.1') == 1
    testo = ' '.join(t)
    assert voce_1341 in testo
    assert risolutiva in testo


@pytest.mark.parametrize("lingua,inizio", [("it", "a) La Segreteria"), ("en", "a) The Organizing")])
def test_premessa_a_non_in_grassetto(lingua, inizio):
    d = Document(str(CARTELLA / f"template_contratto_sponsor_non_ecm_{lingua}.docx"))
    p = next(p for p in d.paragraphs if p.text.strip().startswith(inizio))
    assert p.runs[0].text == "a) " and p.runs[0].bold
    assert all(not r.bold for r in p.runs[1:])
