"""Cliente ESENTE IVA: preventivo, domanda e promemoria delle scadenze
riportano il motivo dell'esenzione e nessun importo IVA."""
import pytest
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.core import mail
from django.utils import timezone

from catalog.models import Service
from contracts.models import (Contract, ContractLine, ContractKind,
                              ContractStatus, Deadline)
from events.models import Event

MOTIVO = 'NON IMPONIBILE Art. 8/1-c "ESPORTATORE ABITUALE"'


@pytest.fixture
def contratto_esente(db, sponsor):
    event = Event.objects.create(
        name={'it': 'Ev Esente', 'en': 'Ev Exempt'}, code='EVE',
        start_date=date(2027, 3, 4), end_date=date(2027, 3, 6))
    service = Service.objects.create(
        event=event, name={'it': 'Stand', 'en': 'Stand'},
        base_price=Decimal('1000.00'), vat_rate=Decimal('22.00'))
    contract = Contract.objects.create(
        sponsor=sponsor, event=event, contract_kind=ContractKind.MAIN,
        status=ContractStatus.SENT, contract_number='EVE-27-001',
        vat_applicable=False, vat_exemption_reason=MOTIVO,
        deposit_percent=Decimal('40'))
    ContractLine.objects.create(contract=contract, service=service, quantity=1)
    contract.refresh_from_db()
    return contract


def _testo(doc):
    from pypdf import PdfReader
    rel = doc.storage_url.replace(settings.MEDIA_URL, '', 1).lstrip('/')
    pdf = PdfReader(str(Path(settings.MEDIA_ROOT) / rel))
    return ' '.join((p.extract_text() or '') for p in pdf.pages)


@pytest.mark.django_db
def test_preventivo_indica_esenzione(contratto_esente):
    from contracts.services.pdf_generator import generate_quote_pdf_html
    testo = _testo(generate_quote_pdf_html(contratto_esente))
    assert 'Esenzione IVA: NON IMPONIBILE Art. 8/1-c' in testo
    assert '1.000,00' in testo or '1000,00' in testo or '1000.00' in testo


def test_riga_iva_domanda_diventa_esenzione(tmp_path):
    from docx import Document
    from contracts.services.pdf_generator import _rimuovi_riga_iva_domanda
    d = Document()
    t = d.add_table(rows=3, cols=3)
    for r, (a, b) in enumerate([('TOTALE IMPONIBILE', '1.000,00'),
                                ('IVA 22%', '220,00'),
                                ('TOTALE', '1.000,00')]):
        t.cell(r, 1).text = a
        t.cell(r, 2).text = b
    p = tmp_path / 'd.docx'
    d.save(str(p))
    assert _rimuovi_riga_iva_domanda(p, MOTIVO, 'it')
    righe = [[c.text for c in r.cells] for r in Document(str(p)).tables[0].rows]
    assert len(righe) == 3
    assert righe[1][1] == f'ESENZIONE IVA: {MOTIVO}'
    assert righe[1][2] == ''


def test_riga_iva_domanda_senza_motivo_sparisce(tmp_path):
    from docx import Document
    from contracts.services.pdf_generator import _rimuovi_riga_iva_domanda
    d = Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = 'IVA 22%', '220,00'
    t.cell(1, 0).text, t.cell(1, 1).text = 'TOTALE', '1.000,00'
    p = tmp_path / 'd.docx'
    d.save(str(p))
    assert _rimuovi_riga_iva_domanda(p)
    assert len(Document(str(p)).tables[0].rows) == 1


@pytest.mark.django_db
def test_promemoria_acconto_esente(contratto_esente):
    from sponsors.models import Contact, ContactRole
    from contracts.tasks.notifications import send_deadline_reminder
    Contact.objects.create(
        sponsor=contratto_esente.sponsor, full_name='Mario Pagatore',
        email='pagatore@test.it', roles=[ContactRole.OPERATIONAL],
        is_primary=True)
    contratto_esente.status = ContractStatus.SIGNED
    contratto_esente.signed_date = date(2026, 10, 1)
    contratto_esente.save()
    deadline = Deadline.objects.create(
        contract=contratto_esente, deadline_type='pagamento_acconto',
        title='Scadenza acconto',
        due_date=timezone.now().date() + timedelta(days=7))
    mail.outbox.clear()
    send_deadline_reminder(deadline.id, reminder_type='reminder')
    assert mail.outbox
    corpo = ' '.join([mail.outbox[0].body or ''] + [
        a for a, _ in getattr(mail.outbox[0], 'alternatives', [])])
    assert 'IVA inclusa' not in corpo
    assert 'esente IVA' in corpo
    assert '400' in corpo


def test_nota_iva():
    from types import SimpleNamespace
    from contracts.tasks.notifications import nota_iva_importo
    con = SimpleNamespace(vat_applicable=True, vat_exemption_reason='')
    senza = SimpleNamespace(vat_applicable=False, vat_exemption_reason='Art. 8')
    assert nota_iva_importo(con) == '(IVA inclusa)'
    assert nota_iva_importo(senza) == '(esente IVA: Art. 8)'
    assert nota_iva_importo(senza, 'en') == '(VAT exempt: Art. 8)'
