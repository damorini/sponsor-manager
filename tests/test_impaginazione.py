"""Documenti Word: un capitolo (titolo + testo, tabelle comprese) resta unito
e passa intero alla pagina dopo se non ci sta."""
from docx import Document
from docx.oxml.ns import qn


def _kn(p):
    ppr = p._p.find(qn('w:pPr'))
    return ppr is not None and ppr.find(qn('w:keepNext')) is not None


def _grassetto(doc, testo):
    p = doc.add_paragraph()
    p.add_run(testo).bold = True
    return p


def test_contratto_articoli_interi(tmp_path):
    from contracts.services.impaginazione import tieni_capitoli_interi

    doc = Document()
    t1 = _grassetto(doc, '1. OBBLIGAZIONI DELLA SEGRETERIA')
    a = doc.add_paragraph('1.1 Primo comma.')
    b = doc.add_paragraph('1.2 Secondo comma.')
    t2 = _grassetto(doc, '2. OBBLIGHI DELLO SPONSOR')
    sotto = _grassetto(doc, '2.1 RINUNCIA E CANCELLAZIONI')
    c = doc.add_paragraph('Testo finale.')
    f = tmp_path / 'c.docx'
    doc.save(f)

    assert tieni_capitoli_interi(f) == 2
    d = Document(f)
    p = {x.text: x for x in d.paragraphs}
    # dentro l'articolo tutto resta col successivo, l'ultimo comma no
    assert _kn(p[t1.text]) and _kn(p[a.text]) and not _kn(p[b.text])
    # il sottotitolo 2.1 non apre un capitolo nuovo
    assert _kn(p[t2.text]) and _kn(p[sotto.text]) and not _kn(p[c.text])


def test_tabella_con_titolo_non_si_spezza(tmp_path):
    from contracts.services.impaginazione import tieni_capitoli_interi

    doc = Document()
    doc.add_paragraph('Termini di pagamento.')
    tab = doc.add_table(rows=3, cols=2)
    tab.cell(0, 0).text = 'RIEPILOGO TECNICO DELLO SPAZIO'
    tab.cell(1, 0).text = 'Tipologia'
    tab.cell(2, 0).text = 'Misure'
    f = tmp_path / 'd.docx'
    doc.save(f)

    assert tieni_capitoli_interi(f) == 2
    t = Document(f).tables[0]
    righe = t._tbl.findall(qn('w:tr'))
    assert all(r.find(qn('w:trPr')).find(qn('w:cantSplit')) is not None for r in righe)
    assert _kn(t.cell(0, 0).paragraphs[0]) and _kn(t.cell(1, 0).paragraphs[0])
    assert not _kn(t.cell(2, 0).paragraphs[0])
    # il paragrafo prima non e' agganciato alla tabella (capitolo diverso)
    assert not _kn(Document(f).paragraphs[0])
