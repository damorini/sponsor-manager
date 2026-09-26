"""Allegato 2: numero di pagine reale; Domanda: termini di cancellazione come
il 3.3 del contratto."""
from datetime import date
from pathlib import Path

import pytest
from django.conf import settings
from docx import Document

from contracts.models import Contract, ContractKind, ContractStatus
from events.models import Event

CARTELLA = Path(settings.BASE_DIR) / 'contracts' / 'templates_pdf'


@pytest.fixture
def contratto(db, sponsor):
    ev = Event.objects.create(name={'it': 'Ev Pag'}, code='PAG',
                              start_date=date(2027, 2, 25), end_date=date(2027, 2, 27))
    return Contract.objects.create(sponsor=sponsor, event=ev, contract_kind=ContractKind.MAIN,
                                   status=ContractStatus.SIGNED, contract_number='PAG-1')


def test_segnaposto_pagine(contratto, tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    from contracts.services.allegato2 import prepara_docx, usa_pagine_allegato
    d0 = Document()
    d0.add_paragraph('Il presente documento si compone di N. {{ pagine_allegato }} pagine.')
    percorso = tmp_path / 'a2.docx'
    d0.save(str(percorso))
    assert usa_pagine_allegato(percorso)
    prepara_docx(percorso, contratto, pagine=6)
    testo = '\n'.join(p.text for p in Document(str(percorso)).paragraphs)
    assert 'si compone di N. 6 pagine.' in testo
    assert not usa_pagine_allegato(percorso)


@pytest.mark.parametrize('lingua,attesi', [
    ('it', ['art. 1385, secondo comma', 'art. 1382 del Codice Civile', 'pari al 100%',
            'concordate per iscritto']),
    ('en', ['art. 1385, second paragraph', 'art. 1382 of the Italian', 'equal to 100%',
            'agreed in writing']),
])
def test_domanda_termini_cancellazione(lingua, attesi):
    t = [p.text for p in Document(
        str(CARTELLA / f'template_domanda_ammissione_{lingua}.docx')).paragraphs]
    i = next(i for i, x in enumerate(t) if x.strip() in ('TERMINI DI CANCELLAZIONE',
                                                            'CANCELLATION TERMS'))
    for a in attesi:
        assert a in t[i + 1], a
