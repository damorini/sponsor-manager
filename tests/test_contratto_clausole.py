"""Contratto di sponsorizzazione (IT): art. 3 con termini di cancellazione e
3.3 (penale = caparra, come nella Domanda), 7.5-7.7, art. 8 e 9, pagine
dichiarate, dichiarazione ex artt. 1341-1342 c.c. con doppia firma."""
import re
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from django.conf import settings
from docx import Document

CARTELLA = Path(settings.BASE_DIR) / 'contracts' / 'templates_pdf'
MODELLI = {
    'it': ('template_contratto_sponsor_non_ecm_it.docx',
           '3. Corrispettivo, modalità di pagamento e termini di cancellazione',
           'articoli 1341 e 1342 del Codice Civile', 'Segreteria Organizzativa'),
    'en': ('template_contratto_sponsor_non_ecm_en.docx',
           '3. Consideration, payment terms and cancellation terms',
           'articles 1341 and 1342 of the Italian Civil Code', 'Organizing Secretariat'),
}


@pytest.mark.parametrize('lingua', ['it', 'en'])
def test_numerazione_e_articoli_nuovi(lingua):
    nome, titolo, art1341, _firma = MODELLI[lingua]
    t = [p.text.strip() for p in Document(str(CARTELLA / nome)).paragraphs]
    assert titolo in t
    inizi = [x.split(' ')[0] for x in t if re.match(r'^\d+\.\d', x)]
    for n in ('3.3', '4.1', '4.2', '8.4', '8.5', '8.6', '12.1', '13.1'):
        assert n in inizi, n
    assert inizi.count('8.4') == 1
    assert '8.7' not in inizi   # l'8.3 sul foro e' stato tolto (resta il 13.1)
    testo = ' '.join(t)
    assert '{{ penale_percent }}%' in testo
    assert '{{ numero_pagine }} pag' in testo
    assert art1341 in testo


@pytest.mark.parametrize('lingua', ['it', 'en'])
def test_doppia_firma(lingua):
    nome, _t, _a, firma = MODELLI[lingua]
    tabelle = Document(str(CARTELLA / nome)).tables
    firme = [tb for tb in tabelle if firma in tb.rows[0].cells[0].text]
    assert len(firme) == 2


@pytest.mark.django_db
def test_penale_uguale_alla_caparra(sponsor):
    from events.models import Event
    from contracts.models import Contract, ContractKind, ContractStatus
    from contracts.services.pdf_generator import _penale_cancellazione
    ev = Event.objects.create(name={'it': 'Pen'}, code='PEN',
                              start_date=date(2027, 2, 1), end_date=date(2027, 2, 2))
    c = Contract.objects.create(sponsor=sponsor, event=ev,
                                contract_kind=ContractKind.MAIN,
                                status=ContractStatus.SENT,
                                deposit_percent=Decimal('45'))
    perc, caparra = _penale_cancellazione(c)
    if caparra:
        assert perc == '45'
    else:
        assert perc == str(ev.cancellation_penalty_percent or 50)
