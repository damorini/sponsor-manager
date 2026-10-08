"""Domanda di ammissione: sotto ogni voce la sua descrizione (lo stand con
l'elenco di cio' che include, i badge con cosa consentono) e i servizi
inclusi con la dicitura 'Incluso' al posto dei prezzi, come nel preventivo."""
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from django.conf import settings
from docx import Document as DocxDocument


@pytest.mark.django_db
def test_descrizioni_e_inclusi(sponsor):
    from sponsors.models import Contact, ContactRole
    from events.models import Event
    from catalog.models import Service, ServiceInclusion
    from contracts.models import Contract, ContractKind, ContractStatus, ContractLine
    from contracts.services.pdf_generator import generate_admission_request_pdf

    Contact.objects.create(
        sponsor=sponsor, full_name='Firma Tario', email='firma@test.it',
        roles=[ContactRole.OPERATIONAL], is_signer=True)
    event = Event.objects.create(
        name={'it': 'Incl Ev', 'en': 'Incl Ev'}, code='IN',
        start_date=date(2026, 11, 1), end_date=date(2026, 11, 2))
    contract = Contract.objects.create(
        sponsor=sponsor, event=event, contract_kind=ContractKind.MAIN,
        status=ContractStatus.SENT, contract_number='IN-26-001')
    stand = Service.objects.create(
        event=event, code='STAND_X', name={'it': 'Spazio espositivo', 'en': 'Space'},
        base_price=Decimal('1000.00'))
    badge = Service.objects.create(
        event=event, code='BADGE_FULL', name={'it': 'Badge Full', 'en': 'Badge Full'},
        description={'it': "Consente l'accesso alle aree espositive"},
        base_price=Decimal('50.00'))
    ServiceInclusion.objects.create(parent=stand, child=badge, quantity=2)
    ContractLine.objects.create(
        contract=contract, service=stand, quantity=1,
        service_description_snapshot='· Area nuda\n· Logo su sito')

    generate_admission_request_pdf(contract)
    docx_path = (Path(settings.MEDIA_ROOT) / 'documents' / 'contracts'
                 / str(contract.id)
                 / f'domanda_ammissione_{contract.contract_number}_{event.id}.docx')
    d = DocxDocument(str(docx_path))
    tab = next(t for t in d.tables
               if 'descrizione dei servizi' in ' '.join(
                   c.text.lower() for c in t.rows[0].cells))
    righe = {r.cells[1].paragraphs[0].text.strip(): r
             for r in tab.rows[1:] if r.cells[0].text.strip()}

    r_stand = next(r for k, r in righe.items() if k.startswith('Spazio espositivo'))
    testi = [p.text for p in r_stand.cells[1].paragraphs]
    assert testi[1:] == ['•  Area nuda', '•  Logo su sito']   # elenco puntato
    assert r_stand.cells[3].text.strip() != 'Incluso'

    r_badge = next(r for k, r in righe.items() if k.startswith('Badge Full'))
    assert "Consente l'accesso alle aree espositive" in r_badge.cells[1].text
    assert r_badge.cells[2].text.strip() == 'Incluso'
    assert r_badge.cells[3].text.strip() == 'Incluso'
    assert '50,00' not in r_badge.cells[2].text
